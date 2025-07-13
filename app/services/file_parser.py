"""
File parsing service for TailerAI v2.0
Modular, reusable document parsing with comprehensive error handling
"""

import logging
import tempfile
from pathlib import Path
from typing import Dict, Any, Optional, List
from dataclasses import dataclass
import re

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    from docx import Document
except ImportError:
    Document = None

from app.config.settings import get_settings

logger = logging.getLogger(__name__)
settings = get_settings()


@dataclass
class ParsedContent:
    """Structured representation of parsed document content"""
    raw_text: str
    contact_info: Dict[str, str]
    sections: Dict[str, List[str]]
    metadata: Dict[str, Any]
    parsing_errors: List[str]


class FileParsingError(Exception):
    """Custom exception for file parsing errors"""
    pass


class DocumentParser:
    """
    Modular document parser supporting multiple formats
    Following best practices: separate concerns, error handling, logging
    """
    
    SUPPORTED_EXTENSIONS = {'.pdf', '.docx', '.txt'}
    
    def __init__(self):
        self.parsers = {
            '.pdf': self._parse_pdf,
            '.docx': self._parse_docx,
            '.txt': self._parse_txt
        }
    
    def parse_document(self, file_path: Path) -> ParsedContent:
        """
        Main parsing entry point with comprehensive error handling
        
        Args:
            file_path: Path to the document file
            
        Returns:
            ParsedContent: Structured parsed content
            
        Raises:
            FileParsingError: When parsing fails
        """
        logger.info(f"Starting document parsing: {file_path}")
        
        try:
            # Validate file existence and extension
            if not file_path.exists():
                raise FileParsingError(f"File not found: {file_path}")
            
            extension = file_path.suffix.lower()
            if extension not in self.SUPPORTED_EXTENSIONS:
                raise FileParsingError(f"Unsupported file type: {extension}")
            
            # Check for required libraries
            self._validate_parser_dependencies(extension)
            
            # Parse document using appropriate parser
            parser_func = self.parsers[extension]
            raw_text = parser_func(file_path)
            
            # Extract structured content
            parsed_content = self._extract_structured_content(raw_text)
            
            logger.info(f"Successfully parsed document: {len(raw_text)} characters extracted")
            return parsed_content
            
        except Exception as e:
            logger.error(f"Document parsing failed: {str(e)}")
            if isinstance(e, FileParsingError):
                raise
            raise FileParsingError(f"Unexpected parsing error: {str(e)}")
    
    def _validate_parser_dependencies(self, extension: str) -> None:
        """Validate required libraries are available"""
        if extension == '.pdf' and PyPDF2 is None:
            raise FileParsingError("PyPDF2 library not available for PDF parsing")
        if extension == '.docx' and Document is None:
            raise FileParsingError("python-docx library not available for DOCX parsing")
    
    def _parse_pdf(self, file_path: Path) -> str:
        """
        Parse PDF file using PyPDF2
        Opens in binary mode as per best practices
        """
        logger.debug(f"Parsing PDF: {file_path}")
        
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                
                if len(pdf_reader.pages) == 0:
                    raise FileParsingError("PDF contains no pages")
                
                # Extract text from all pages
                text_content = []
                for page_num, page in enumerate(pdf_reader.pages):
                    try:
                        page_text = page.extract_text()
                        if page_text.strip():
                            text_content.append(page_text)
                        logger.debug(f"Extracted {len(page_text)} characters from page {page_num + 1}")
                    except Exception as e:
                        logger.warning(f"Failed to extract text from page {page_num + 1}: {e}")
                
                if not text_content:
                    raise FileParsingError("No readable text found in PDF")
                
                return '\n'.join(text_content)
                
        except PyPDF2.errors.PdfReadError as e:
            raise FileParsingError(f"Invalid or corrupted PDF file: {e}")
        except Exception as e:
            raise FileParsingError(f"PDF parsing error: {e}")
    
    def _parse_docx(self, file_path: Path) -> str:
        """
        Parse DOCX file using python-docx
        Extracts text from paragraphs and tables
        """
        logger.debug(f"Parsing DOCX: {file_path}")
        
        try:
            doc = Document(file_path)
            text_content = []
            
            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text_content.append(' | '.join(row_text))
            
            if not text_content:
                raise FileParsingError("No readable text found in DOCX")
            
            return '\n'.join(text_content)
            
        except Exception as e:
            raise FileParsingError(f"DOCX parsing error: {e}")
    
    def _parse_txt(self, file_path: Path) -> str:
        """
        Parse plain text file
        Opens in text mode as per best practices
        """
        logger.debug(f"Parsing TXT: {file_path}")
        
        try:
            # Try UTF-8 first, fallback to latin-1 for compatibility
            for encoding in ['utf-8', 'latin-1']:
                try:
                    with open(file_path, 'r', encoding=encoding) as file:
                        content = file.read()
                        if content.strip():
                            return content
                        raise FileParsingError("Text file is empty")
                except UnicodeDecodeError:
                    if encoding == 'latin-1':  # Last attempt
                        raise FileParsingError("Unable to decode text file")
                    continue
                    
        except Exception as e:
            raise FileParsingError(f"Text file parsing error: {e}")
    
    def _extract_structured_content(self, raw_text: str) -> ParsedContent:
        """
        Extract structured content from raw text
        Modular approach with separate extraction methods
        """
        parsing_errors = []
        
        try:
            contact_info = self._extract_contact_info(raw_text)
        except Exception as e:
            logger.warning(f"Contact info extraction failed: {e}")
            contact_info = {}
            parsing_errors.append(f"Contact extraction error: {e}")
        
        try:
            sections = self._extract_sections(raw_text)
        except Exception as e:
            logger.warning(f"Section extraction failed: {e}")
            sections = {"content": [raw_text]}  # Fallback
            parsing_errors.append(f"Section extraction error: {e}")
        
        metadata = {
            "character_count": len(raw_text),
            "word_count": len(raw_text.split()),
            "line_count": len(raw_text.splitlines())
        }
        
        return ParsedContent(
            raw_text=raw_text,
            contact_info=contact_info,
            sections=sections,
            metadata=metadata,
            parsing_errors=parsing_errors
        )
    
    def _extract_contact_info(self, text: str) -> Dict[str, str]:
        """Extract contact information using regex patterns"""
        contact_info = {}
        
        # Email pattern
        email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
        email_match = re.search(email_pattern, text)
        if email_match:
            contact_info['email'] = email_match.group()
        
        # Phone pattern (flexible formats)
        phone_pattern = r'(\+?1[-.\s]?)?\(?([0-9]{3})\)?[-.\s]?([0-9]{3})[-.\s]?([0-9]{4})'
        phone_match = re.search(phone_pattern, text)
        if phone_match:
            contact_info['phone'] = phone_match.group()
        
        # LinkedIn pattern
        linkedin_pattern = r'linkedin\.com/in/[\w-]+'
        linkedin_match = re.search(linkedin_pattern, text, re.IGNORECASE)
        if linkedin_match:
            contact_info['linkedin'] = linkedin_match.group()
        
        # Extract name (first line heuristic)
        lines = text.strip().split('\n')
        if lines:
            first_line = lines[0].strip()
            # Simple heuristic: if first line looks like a name
            if len(first_line.split()) <= 4 and not any(char.isdigit() for char in first_line):
                contact_info['name'] = first_line
        
        return contact_info
    
    def _extract_sections(self, text: str) -> Dict[str, List[str]]:
        """
        Extract resume sections using common patterns
        Modular approach for easy extension
        """
        sections = {}
        
        # Common section headers
        section_patterns = {
            'education': r'(EDUCATION|Education|ACADEMIC|Academic)',
            'experience': r'(EXPERIENCE|Experience|WORK|Work|EMPLOYMENT|Employment)',
            'skills': r'(SKILLS|Skills|TECHNICAL|Technical|COMPETENCIES|Competencies)',
            'projects': r'(PROJECTS|Projects|PROJECT|Project)',
            'certifications': r'(CERTIFICATIONS|Certifications|CERTIFICATES|Certificates)',
            'awards': r'(AWARDS|Awards|HONORS|Honors|ACHIEVEMENTS|Achievements)'
        }
        
        lines = text.split('\n')
        current_section = 'header'
        sections[current_section] = []
        
        for line in lines:
            line = line.strip()
            if not line:
                continue
            
            # Check if line matches any section header
            section_found = False
            for section_name, pattern in section_patterns.items():
                if re.match(pattern, line, re.IGNORECASE):
                    current_section = section_name
                    sections[current_section] = []
                    section_found = True
                    break
            
            if not section_found:
                sections[current_section].append(line)
        
        # Remove empty sections
        sections = {k: v for k, v in sections.items() if v}
        
        return sections


# Factory function for dependency injection
def create_document_parser() -> DocumentParser:
    """Factory function for creating document parser instance"""
    return DocumentParser()